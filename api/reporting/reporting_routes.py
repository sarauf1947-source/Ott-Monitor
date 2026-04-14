# -*- coding: utf-8 -*-
"""Reporting endpoints - JSON summary, CSV, XLSX, HTML export. api/reporting/reporting_routes.py"""
import csv
import io
import uuid
from datetime import datetime, timedelta
from typing import List

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from sqlalchemy import func, select
from sqlalchemy.ext.asyncio import AsyncSession

from api.auth.auth_deps import get_current_user
from db.database import get_db

reporting_router = APIRouter(prefix="/api/v1/reports", tags=["Reporting"])


async def _build_report(db: AsyncSession, channel_id: uuid.UUID,
                        start: datetime, end: datetime):
    from db.models import Channel, Metric, Error

    ch = (await db.execute(select(Channel).where(Channel.id == channel_id))).scalar_one_or_none()
    if not ch:
        return None

    # Metric aggregates - uses actual column names from models.py
    m = (await db.execute(
        select(
            func.count(Metric.id).label("total"),
            func.avg(Metric.bitrate).label("avg_bitrate"),
            func.min(Metric.bitrate).label("min_bitrate"),
            func.max(Metric.bitrate).label("max_bitrate"),
            func.avg(Metric.response_time).label("avg_response"),
        ).where(Metric.channel_id == channel_id,
                Metric.timestamp.between(start, end))
    )).one()

    up = (await db.execute(
        select(func.count(Metric.id)).where(
            Metric.channel_id == channel_id,
            Metric.timestamp.between(start, end),
            Metric.is_available == True,  # noqa: E712
        )
    )).scalar_one() or 0

    total    = m.total or 1
    uptime   = round((up / total) * 100, 2)

    errs = (await db.execute(
        select(Error.error_type, Error.severity, func.count(Error.id).label("count"))
        .where(Error.channel_id == channel_id, Error.timestamp.between(start, end))
        .group_by(Error.error_type, Error.severity)
        .order_by(func.count(Error.id).desc())
    )).fetchall()

    def val(v):
        return v.value if hasattr(v, "value") else str(v)

    return {
        "channel_id":     str(ch.id),
        "channel_name":   ch.name,
        "stream_url":     ch.stream_url,
        "current_status": val(ch.status),
        "report_period":  {"start": start.isoformat(), "end": end.isoformat()},
        "uptime":  {"percent": uptime, "total_checks": total, "up_checks": up},
        "metrics": {
            "avg_bitrate_kbps": round(m.avg_bitrate or 0, 1),
            "min_bitrate_kbps": round(m.min_bitrate or 0, 1),
            "max_bitrate_kbps": round(m.max_bitrate or 0, 1),
            "avg_response_ms":  round(m.avg_response or 0, 1),
        },
        "errors": [{"type": val(r.error_type), "severity": val(r.severity), "count": r.count}
                   for r in errs],
    }


@reporting_router.get("/channels/{channel_id}/summary")
async def channel_summary(
    channel_id: uuid.UUID,
    days: int = Query(7, ge=1, le=90),
    db: AsyncSession = Depends(get_db),
    _u=Depends(get_current_user),
):
    end  = datetime.utcnow()
    data = await _build_report(db, channel_id, end - timedelta(days=days), end)
    if not data:
        raise HTTPException(404, "Channel not found")
    return data


@reporting_router.get("/channels/{channel_id}/export/csv")
async def export_csv(
    channel_id: uuid.UUID,
    days: int = Query(7, ge=1, le=90),
    db: AsyncSession = Depends(get_db),
    _u=Depends(get_current_user),
):
    end  = datetime.utcnow()
    data = await _build_report(db, channel_id, end - timedelta(days=days), end)
    if not data:
        raise HTTPException(404, "Channel not found")

    buf = io.StringIO()
    w   = csv.writer(buf)
    w.writerow(["OTT Monitor Channel Report"])
    w.writerow(["Channel", data["channel_name"]])
    w.writerow(["Period", f"{data['report_period']['start']} to {data['report_period']['end']}"])
    w.writerow([])
    w.writerow(["Uptime %", "Total Checks", "Up Checks"])
    u = data["uptime"]
    w.writerow([u["percent"], u["total_checks"], u["up_checks"]])
    w.writerow([])
    m = data["metrics"]
    w.writerow(["Avg Bitrate kbps", "Min Bitrate", "Max Bitrate", "Avg Response ms"])
    w.writerow([m["avg_bitrate_kbps"], m["min_bitrate_kbps"], m["max_bitrate_kbps"], m["avg_response_ms"]])
    w.writerow([])
    w.writerow(["Error Type", "Severity", "Count"])
    for e in data["errors"]:
        w.writerow([e["type"], e["severity"], e["count"]])

    fname = f"report_{data['channel_name'].replace(' ', '_')}_{datetime.utcnow().strftime('%Y%m%d')}.csv"
    return StreamingResponse(
        io.BytesIO(buf.getvalue().encode("utf-8")),
        media_type="text/csv",
        headers={"Content-Disposition": f"attachment; filename={fname}"},
    )


@reporting_router.get("/channels/{channel_id}/export/xlsx")
async def export_xlsx(
    channel_id: uuid.UUID,
    days: int = Query(7, ge=1, le=90),
    db: AsyncSession = Depends(get_db),
    _u=Depends(get_current_user),
):
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment
    except ImportError:
        raise HTTPException(500, "openpyxl not installed - add to requirements.txt and rebuild")

    end  = datetime.utcnow()
    data = await _build_report(db, channel_id, end - timedelta(days=days), end)
    if not data:
        raise HTTPException(404, "Channel not found")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Report"
    hf = Font(bold=True, color="FFFFFF")
    hb = PatternFill(fill_type="solid", fgColor="1E3A5F")

    def hrow(vals, row):
        for c, v in enumerate(vals, 1):
            cell = ws.cell(row=row, column=c, value=v)
            cell.font = hf
            cell.fill = hb
            cell.alignment = Alignment(horizontal="center")

    ws["A1"] = f"OTT Monitor - {data['channel_name']}"
    ws["A1"].font = Font(bold=True, size=13)
    ws["A2"] = f"Period: {data['report_period']['start']} to {data['report_period']['end']}"

    hrow(["Uptime %", "Total Checks", "Up Checks"], 4)
    u = data["uptime"]
    ws.append([u["percent"], u["total_checks"], u["up_checks"]])
    ws.append([])
    m = data["metrics"]
    hrow(["Avg Bitrate kbps", "Min Bitrate", "Max Bitrate", "Avg Response ms"], ws.max_row + 1)
    ws.append([m["avg_bitrate_kbps"], m["min_bitrate_kbps"], m["max_bitrate_kbps"], m["avg_response_ms"]])
    ws.append([])
    hrow(["Error Type", "Severity", "Count"], ws.max_row + 1)
    for e in data["errors"]:
        ws.append([e["type"], e["severity"], e["count"]])

    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = min(
            max((len(str(c.value or "")) for c in col), default=8) + 4, 40
        )

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    fname = f"report_{data['channel_name'].replace(' ', '_')}_{datetime.utcnow().strftime('%Y%m%d')}.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={fname}"},
    )


@reporting_router.get("/channels/{channel_id}/export/pdf")
async def export_html(
    channel_id: uuid.UUID,
    days: int = Query(7, ge=1, le=90),
    db: AsyncSession = Depends(get_db),
    _u=Depends(get_current_user),
):
    """Returns printable HTML. Use browser Print > Save as PDF."""
    end  = datetime.utcnow()
    data = await _build_report(db, channel_id, end - timedelta(days=days), end)
    if not data:
        raise HTTPException(404, "Channel not found")

    def sev_color(s):
        return {"CRITICAL": "#dc2626", "MAJOR": "#d97706", "WARNING": "#7c3aed"}.get(s, "#374151")

    rows_html = []
    for e in data["errors"]:
        color = sev_color(e["severity"])
        rows_html.append(
            "<tr><td>" + e["type"] + "</td>"
            "<td style='color:" + color + ";font-weight:600'>" + e["severity"] + "</td>"
            "<td>" + str(e["count"]) + "</td></tr>"
        )
    errors_html = "".join(rows_html) if rows_html else "<tr><td colspan='3' style='color:#9ca3af'>No errors in this period</td></tr>"

    m  = data["metrics"]
    u  = data["uptime"]
    html = (
        "<!DOCTYPE html><html><head><meta charset='utf-8'>"
        f"<title>Report - {data['channel_name']}</title>"
        "<style>"
        "body{font-family:Arial,sans-serif;margin:40px;color:#111827}"
        "h1{color:#1E3A5F;border-bottom:3px solid #1E3A5F;padding-bottom:8px;font-size:22px}"
        "h2{color:#1E3A5F;font-size:16px;margin:24px 0 8px}"
        "table{border-collapse:collapse;width:100%}"
        "th{background:#1E3A5F;color:#fff;padding:8px 14px;text-align:left;font-size:13px}"
        "td{padding:7px 14px;border-bottom:1px solid #e5e7eb;font-size:13px}"
        "tr:nth-child(even){background:#f9fafb}"
        ".kpi{display:inline-block;background:#eff6ff;border-radius:8px;padding:12px 18px;"
        "margin:6px;text-align:center;min-width:110px;border:1px solid #bfdbfe}"
        ".kv{font-size:22px;font-weight:700;color:#1E3A5F}"
        ".kl{font-size:11px;color:#6b7280;margin-top:2px}"
        "@media print{body{margin:20px}}"
        "</style></head><body>"
        "<h1>OTT Monitor - Channel Report</h1>"
        f"<p style='font-size:13px;color:#6b7280'>"
        f"<b>{data['channel_name']}</b> | {data['current_status']}"
        f" | {data['report_period']['start'][:10]} to {data['report_period']['end'][:10]}</p>"
        "<h2>Key Metrics</h2><div>"
        f"<div class='kpi'><div class='kv'>{u['percent']}%</div><div class='kl'>Uptime</div></div>"
        f"<div class='kpi'><div class='kv'>{m['avg_bitrate_kbps']}</div><div class='kl'>Avg Bitrate kbps</div></div>"
        f"<div class='kpi'><div class='kv'>{m['min_bitrate_kbps']}</div><div class='kl'>Min Bitrate</div></div>"
        f"<div class='kpi'><div class='kv'>{m['max_bitrate_kbps']}</div><div class='kl'>Max Bitrate</div></div>"
        f"<div class='kpi'><div class='kv'>{m['avg_response_ms']} ms</div><div class='kl'>Avg Response</div></div>"
        f"<div class='kpi'><div class='kv'>{u['total_checks']}</div><div class='kl'>Total Checks</div></div>"
        "</div>"
        "<h2>Error Breakdown</h2>"
        "<table><thead><tr><th>Error Type</th><th>Severity</th><th>Count</th></tr></thead>"
        f"<tbody>{errors_html}</tbody></table>"
        f"<p style='margin-top:40px;font-size:11px;color:#9ca3af'>"
        f"Generated {datetime.utcnow().strftime('%Y-%m-%d %H:%M')} UTC"
        " - Use browser Print to save as PDF</p>"
        "</body></html>"
    )

    fname = f"report_{data['channel_name'].replace(' ', '_')}_{datetime.utcnow().strftime('%Y%m%d')}.html"
    return StreamingResponse(
        io.BytesIO(html.encode("utf-8")),
        media_type="text/html",
        headers={"Content-Disposition": f"inline; filename={fname}"},
    )


@reporting_router.post("/bulk")
async def bulk_report(
    channel_ids: List[str],
    days: int = Query(7, ge=1, le=90),
    db: AsyncSession = Depends(get_db),
    _u=Depends(get_current_user),
):
    end = datetime.utcnow()
    start = end - timedelta(days=days)
    results = []
    for cid_str in channel_ids:
        try:
            cid = uuid.UUID(cid_str)
        except ValueError:
            continue
        data = await _build_report(db, cid, start, end)
        if data:
            results.append(data)
    return {"report_count": len(results), "reports": results}
